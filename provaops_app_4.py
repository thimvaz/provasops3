"""
ProvaOps v4 — Acelerador (Word -> LaTeX) + Embaralhador (versões B, C, ...).

Base: v3. Principais correções em relação à v3:
  Conversor
    - Reconhece "QUESTÃO 05 –", "Questão 3:" e alternativas "(A)" / "A)" / "a.".
    - Número com ponto decimal ("3.5 m/s", "10.000") não vira mais questão/alternativa.
    - Escapa caracteres especiais do LaTeX (% & # _ { } ~ ^ \\ $ e "R$"), que quebravam a compilação
      ou "comentavam" o resto da linha (ex.: "50% do trajeto").
    - Mantém negrito/itálico/sublinhado, sobrescrito/subscrito (m/s², H₂O) e converte equações do Word.
    - Lê tabelas, hiperlinks e listas automáticas do Word/Google Docs (marcadores a), b), 1., 2. que não são texto).
    - Imagem no mesmo parágrafo do enunciado agora vem DEPOIS do "Questão N" (antes ia parar na questão anterior).
    - Texto de apoio logo após a última alternativa não gruda mais nela.
    - Extensões de imagem corretas (emf/wmf/gif etc.) com aviso quando o pdfLaTeX não aceita.
    - Relatório de conferência: questões detectadas, nº de alternativas, avisos.
  Embaralhador
    - Cabeçalho da prova e texto introdutório da disciplina não somem / não viajam mais com a 1ª questão.
    - Questão discursiva (sem alternativas) não gera mais \\begin{enumerate} vazio (erro de compilação).
    - \\hl{...} com chaves aninhadas ($x^{2}$) removido corretamente (antes estragava o LaTeX).
    - "Todas as anteriores"/"Nenhuma das anteriores" ficam fixas no fim.
    - Aceita "% INÍCIO BLOCO" (com acento) e avisa quando FIM/INÍCIO está sem par.
    - Alternativas nunca saem na mesma ordem da original; versões diferentes entre si.
    - Gabarito mestre (A + todas as versões), CSV em UTF-8 com ';' (abre certo no Excel pt-BR).
    - Avisos de referências a "questão N" no texto (a numeração muda nas versões).
  Interface
    - Resultados ficam em st.session_state (não somem ao clicar em Baixar).
"""

import hashlib
import random
import re
import string
import zipfile
from io import BytesIO

import docx
import pandas as pd
import streamlit as st
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _m(tag):
    return "{%s}%s" % (M_NS, tag)


# ==========================================
# TEXTO: ESCAPE / SÍMBOLOS
# ==========================================

_ESC = {
    "\\": r"\textbackslash{}", "{": r"\{", "}": r"\}", "%": r"\%", "&": r"\&",
    "#": r"\#", "_": r"\_", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
}

_GREGAS = {
    "α": "alpha", "β": "beta", "γ": "gamma", "δ": "delta", "ε": "varepsilon", "ζ": "zeta",
    "η": "eta", "θ": "theta", "ι": "iota", "κ": "kappa", "λ": "lambda", "μ": "mu", "µ": "mu",
    "ν": "nu", "ξ": "xi", "π": "pi", "ρ": "rho", "σ": "sigma", "τ": "tau", "υ": "upsilon",
    "φ": "varphi", "χ": "chi", "ψ": "psi", "ω": "omega",
    "Γ": "Gamma", "Δ": "Delta", "Θ": "Theta", "Λ": "Lambda", "Ξ": "Xi", "Π": "Pi",
    "Σ": "Sigma", "Φ": "Phi", "Ψ": "Psi", "Ω": "Omega",
}
_SIMBOLOS = {k: "\\" + v for k, v in _GREGAS.items()}
_SIMBOLOS.update({
    "−": "-", "×": r"\times", "·": r"\cdot", "±": r"\pm", "∓": r"\mp", "≤": r"\leq",
    "≥": r"\geq", "≠": r"\neq", "≈": r"\approx", "≅": r"\cong", "≡": r"\equiv",
    "∝": r"\propto", "∞": r"\infty", "√": r"\surd", "→": r"\rightarrow", "←": r"\leftarrow",
    "↔": r"\leftrightarrow", "⇒": r"\Rightarrow", "⇔": r"\Leftrightarrow", "∑": r"\sum",
    "∏": r"\prod", "∫": r"\int", "∂": r"\partial", "∇": r"\nabla", "∈": r"\in",
    "∅": r"\emptyset",
})
_SUPERS = {"⁰": "0", "⁴": "4", "⁵": "5", "⁶": "6", "⁷": "7", "⁸": "8", "⁹": "9", "⁻": "-", "⁺": "+"}
_SUBS = {"₀": "0", "₁": "1", "₂": "2", "₃": "3", "₄": "4", "₅": "5", "₆": "6", "₇": "7", "₈": "8", "₉": "9"}


def _escapar(texto, estado):
    """Escapa `texto` para LaTeX. `estado` carrega {math_ok, em_math, prev} entre trechos."""
    out = []
    for c in texto:
        em_math = estado["em_math"]
        if c == "$":
            if estado["math_ok"] and not (estado["prev"] in ("R", "S") and not em_math):
                estado["em_math"] = not em_math
                out.append("$")
            else:
                out.append(r"\$")
        elif c == "\n":
            out.append(" " if em_math else r" \\ ")
        elif c == "\t":
            out.append(" ")
        elif c in _SIMBOLOS:
            out.append(_SIMBOLOS[c] + " " if em_math else r"\ensuremath{%s}" % _SIMBOLOS[c])
        elif c in _SUPERS:
            out.append("^{%s}" % _SUPERS[c] if em_math else r"\textsuperscript{%s}" % _SUPERS[c])
        elif c in _SUBS:
            out.append("_{%s}" % _SUBS[c] if em_math else r"\textsubscript{%s}" % _SUBS[c])
        elif em_math:
            out.append({"%": r"\%", "&": r"\&", "#": r"\#"}.get(c, c))
        else:
            out.append(_ESC.get(c, c))
        estado["prev"] = c
    return "".join(out)


def _mesclar(segs):
    """Junta trechos vizinhos com a mesma formatação (o Word fragmenta runs à toa)."""
    out = []
    for s in segs:
        if out and not s["eq"] and not out[-1]["eq"] and all(
                s[k] == out[-1][k] for k in ("b", "i", "u", "sup", "sub")):
            out[-1] = dict(out[-1], t=out[-1]["t"] + s["t"])
        else:
            out.append(dict(s))
    return out


def _segmentos_para_latex(segs):
    plano = "".join(s["t"] for s in segs if not s["eq"])
    n_dolar = len(re.findall(r"(?<![RS])\$", plano))
    math_ok = n_dolar > 0 and n_dolar % 2 == 0
    estado = {"math_ok": math_ok, "em_math": False, "prev": ""}
    partes = []
    for s in _mesclar(segs):
        if s["eq"]:
            partes.append("$" + s["t"] + "$")
            continue
        txt = _escapar(s["t"], estado)
        if not math_ok and txt.strip():
            esq = txt[: len(txt) - len(txt.lstrip())]
            dir_ = txt[len(txt.rstrip()):]
            nucleo = txt.strip()
            if s["sup"]:
                nucleo = r"\textsuperscript{%s}" % nucleo
            if s["sub"]:
                nucleo = r"\textsubscript{%s}" % nucleo
            if s["b"]:
                nucleo = r"\textbf{%s}" % nucleo
            if s["i"]:
                nucleo = r"\textit{%s}" % nucleo
            if s["u"]:
                nucleo = r"\underline{%s}" % nucleo
            txt = esq + nucleo + dir_
        partes.append(txt)
    out = "".join(partes).strip()
    return re.sub(r"^(?:\s*\\\\\s*)+|(?:\s*\\\\\s*)+$", "", out).strip()


def _cortar_prefixo(segs, n):
    """Remove os primeiros n caracteres (de texto plano) da lista de trechos."""
    out, resto = [], n
    for s in segs:
        tam = len(s["t"])
        if resto >= tam:
            resto -= tam
            continue
        if resto > 0:
            s = dict(s, t=s["t"][resto:])
            resto = 0
        out.append(s)
    return out


# ==========================================
# WORD: EQUAÇÕES (OMML -> LaTeX), LISTAS, PARÁGRAFOS
# ==========================================

_DELIMS = {"(": r"\left(", ")": r"\right)", "[": r"\left[", "]": r"\right]",
           "{": r"\left\{", "}": r"\right\}", "|": r"\left|", "": None}
_DELIMS_FIM = {"(": r"\right)", ")": r"\right)", "[": r"\right]", "]": r"\right]",
               "{": r"\right\}", "}": r"\right\}", "|": r"\right|"}
_NARY = {"∑": r"\sum", "∏": r"\prod", "∫": r"\int", "∬": r"\iint", "∮": r"\oint"}


def _omml_texto(el):
    txt = "".join((t.text or "") for t in el if t.tag in (_m("t"), qn("w:t")))
    out = []
    for c in txt:
        if c == "°":
            out.append(r"^{\circ}")
        elif c in _SIMBOLOS:
            out.append(_SIMBOLOS[c] + " ")
        elif c in "%&#":
            out.append("\\" + c)
        else:
            out.append(c)
    return "".join(out)


def omml_para_latex(el):
    """Conversão modesta, mas cobre o que aparece em prova (frações, potências, raízes, parênteses)."""
    tag = el.tag
    if not isinstance(tag, str):
        return ""

    def filhos(e):
        return "".join(omml_para_latex(c) for c in e)

    def sub(nome):
        achado = el.find(_m(nome))
        return filhos(achado) if achado is not None else ""

    if tag == _m("r"):
        return _omml_texto(el)
    if tag.endswith("Pr") and tag.startswith("{" + M_NS):
        return ""
    if tag == _m("f"):
        return r"\frac{%s}{%s}" % (sub("num"), sub("den"))
    if tag == _m("sSup"):
        return "{%s}^{%s}" % (sub("e"), sub("sup"))
    if tag == _m("sSub"):
        return "{%s}_{%s}" % (sub("e"), sub("sub"))
    if tag == _m("sSubSup"):
        return "{%s}_{%s}^{%s}" % (sub("e"), sub("sub"), sub("sup"))
    if tag == _m("rad"):
        grau = sub("deg").strip()
        return (r"\sqrt[%s]{%s}" % (grau, sub("e"))) if grau else (r"\sqrt{%s}" % sub("e"))
    if tag == _m("d"):
        def chr_(nome, padrao):
            c = el.find(_m("dPr") + "/" + _m(nome))
            if c is not None and c.get(_m("val")) is not None:
                return c.get(_m("val"))
            return padrao
        ini, fim, sep = chr_("begChr", "("), chr_("endChr", ")"), chr_("sepChr", "|")
        corpo = (sep if sep else " ").join(filhos(e) for e in el.findall(_m("e")))
        esq = _DELIMS.get(ini, ini) or r"\left."
        dir_ = _DELIMS_FIM.get(fim, r"\right.") if fim else r"\right."
        return f"{esq} {corpo} {dir_}"
    if tag == _m("nary"):
        c = el.find(_m("naryPr") + "/" + _m("chr"))
        op = _NARY.get(c.get(_m("val")) if c is not None else "∫", r"\int")
        inf, sup = sub("sub"), sub("sup")
        return op + (("_{%s}" % inf) if inf else "") + (("^{%s}" % sup) if sup else "") + " " + sub("e")
    return filhos(el)


def _novo_seg(texto, **fmt):
    base = {"t": texto, "b": False, "i": False, "u": False, "sup": False, "sub": False, "eq": False}
    base.update(fmt)
    return base


def _segmentos_do_paragrafo(para):
    """Lista de trechos [{t, b, i, u, sup, sub, eq}] na ordem do XML (inclui hyperlinks e equações)."""
    segs = []
    info = {"equacoes": 0}

    def add_run(r_el):
        run = Run(r_el, para)
        txt = run.text
        if not txt:
            return
        f = run.font
        segs.append(_novo_seg(
            txt.replace("\t", " "),
            b=bool(run.bold), i=bool(run.italic), u=bool(run.underline),
            sup=bool(f.superscript), sub=bool(f.subscript)))

    def add_eq(omath):
        latex = omml_para_latex(omath).strip()
        if latex:
            segs.append(_novo_seg(latex, eq=True))
            info["equacoes"] += 1

    def walk(el):
        for ch in el.iterchildren():
            if ch.tag == qn("w:r"):
                add_run(ch)
            elif ch.tag == _m("oMath"):
                add_eq(ch)
            elif ch.tag == _m("oMathPara"):
                for om in ch.iterchildren(_m("oMath")):
                    add_eq(om)
            elif ch.tag in (qn("w:hyperlink"), qn("w:ins"), qn("w:smartTag"), qn("w:sdt"),
                            qn("w:sdtContent"), qn("w:fldSimple")):
                walk(ch)

    walk(para._p)
    return segs, info["equacoes"]


def _extrair_imagens_do_paragrafo(paragrafo, doc_part):
    """[(bytes, extensao)] na ORDEM em que aparecem no XML."""
    imagens = []
    for blip in paragrafo._element.iter(qn("a:blip")):
        r_id = blip.get(qn("r:embed"))
        if not r_id:
            continue
        try:
            image_part = doc_part.related_parts[r_id]
        except KeyError:
            continue
        try:
            ext = image_part.partname.ext.lower()
        except Exception:
            ext = image_part.content_type.split("/")[-1].lower().replace("x-", "")
        ext = {"jpeg": "jpg"}.get(ext, ext)
        imagens.append((image_part.blob, ext))
    return imagens


def _iterar_paragrafos(doc):
    """(Paragraph, em_tabela) na ordem do documento, incluindo o texto das tabelas."""
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc), False
        elif child.tag == qn("w:tbl"):
            tbl = Table(child, doc)
            for row in tbl.rows:
                vistos = set()
                for cell in row.cells:
                    if id(cell._tc) in vistos:
                        continue
                    vistos.add(id(cell._tc))
                    for p in cell.paragraphs:
                        yield p, True


def _fmt_num(n, fmt):
    if fmt == "lowerLetter" or fmt == "upperLetter":
        s, k = "", n
        while k > 0:
            k, r = divmod(k - 1, 26)
            s = string.ascii_lowercase[r] + s
        return s.upper() if fmt == "upperLetter" else s
    if fmt in ("lowerRoman", "upperRoman"):
        pares = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
                 (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
        s, k = "", n
        for v, r in pares:
            while k >= v:
                s += r
                k -= v
        return s.upper() if fmt == "upperRoman" else s
    if fmt == "decimalZero":
        return f"{n:02d}"
    return str(n)


class Numeracao:
    """Reconstrói o rótulo ("a)", "1.", "(A)") de parágrafos de lista automática do Word/Google Docs."""

    def __init__(self, doc):
        self.abstract, self.num2abs, self.cont = {}, {}, {}
        try:
            raiz = doc.part.numbering_part.element
        except Exception:
            return
        for an in raiz.findall(qn("w:abstractNum")):
            niveis = {}
            for lvl in an.findall(qn("w:lvl")):
                try:
                    il = int(lvl.get(qn("w:ilvl")))
                except (TypeError, ValueError):
                    continue
                st_, fm, tx = lvl.find(qn("w:start")), lvl.find(qn("w:numFmt")), lvl.find(qn("w:lvlText"))
                niveis[il] = (
                    int(st_.get(qn("w:val"))) if st_ is not None else 1,
                    fm.get(qn("w:val")) if fm is not None else "decimal",
                    tx.get(qn("w:val")) if tx is not None else "%1.",
                )
            self.abstract[an.get(qn("w:abstractNumId"))] = niveis
        for n in raiz.findall(qn("w:num")):
            a = n.find(qn("w:abstractNumId"))
            if a is not None:
                self.num2abs[n.get(qn("w:numId"))] = a.get(qn("w:val"))

    @staticmethod
    def _numpr(para):
        try:
            ppr = para._p.pPr
            if ppr is not None and ppr.numPr is not None:
                return ppr.numPr
            est = para.style
            while est is not None:
                ppr = est.element.pPr
                if ppr is not None and ppr.numPr is not None:
                    return ppr.numPr
                est = est.base_style
        except Exception:
            pass
        return None

    def rotulo(self, para):
        np_ = self._numpr(para)
        if np_ is None or np_.numId is None:
            return ""
        nid = str(np_.numId.val)
        il = np_.ilvl.val if np_.ilvl is not None else 0
        if nid == "0":
            return ""
        niveis = self.abstract.get(self.num2abs.get(nid), {})
        if il not in niveis:
            return ""
        inicio, fmt, texto = niveis[il]
        cont = self.cont.setdefault(nid, {})
        if fmt == "bullet":
            return ""
        cont[il] = cont.get(il, inicio - 1) + 1
        for k in [k for k in cont if k > il]:
            del cont[k]

        def num_nivel(m):
            lv = int(m.group(1)) - 1
            ini_lv, fm_lv, _ = niveis.get(lv, (1, "decimal", ""))
            return _fmt_num(cont.get(lv, ini_lv), fm_lv)

        return re.sub(r"%(\d)", num_nivel, texto)


# ==========================================
# ACELERADOR (WORD -> LATEX)
# ==========================================

PREAMBULO_BASE = r"""\documentclass[a4paper,10pt]{exam}
\usepackage[utf8]{inputenc}
\usepackage[T1]{fontenc}
\usepackage[brazil]{babel}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}
\usepackage{soul}
\usepackage[shortlabels]{enumitem}
\usepackage{multicol}

\begin{document}
"""

# "Questão 5", "QUESTÃO 05 –", "Questão 3:" — o rótulo explícito basta como sinal.
RE_Q_FORTE = re.compile(r"^\s*quest[aã]o\s*(\d+)\s*[\.\):\-–—]*\s*", re.IGNORECASE)
# Número solto ("3." / "3)") — sinal fraco, exige espaço depois (para não pegar "3.5" ou "10.000").
RE_Q_FRACA = re.compile(r"^\s*(\d+)[\.\)](?:\s+|$)")
# "a)", "(A)", "b." — exige espaço depois ("e.g." e "a.b" não contam).
RE_ALT = re.compile(r"^\s*(?:\(([a-eA-E])\)|([a-eA-E])([\.\)]))(?:\s+|$)")
RE_HEADING = re.compile(r"^(Heading|T[ií]tulo)\s*\d", re.IGNORECASE)


def _bracket_enum(m):
    """Rótulo do enumerate a partir do 1º marcador da lista: '(a)', 'A)', 'a.' ..."""
    letra = m.group(1) or m.group(2)
    base = "A" if letra.isupper() else "a"
    if m.group(1):
        return f"({base})"
    return f"{base}{m.group(3)}"


def converter_docx_para_latex(docx_file):
    """Retorna (latex, imagens[(bytes, ext)], avisos, notas, resumo[list[dict]])."""
    doc = docx.Document(docx_file)
    doc_part = doc.part
    numeracao = Numeracao(doc)
    avisos, notas, resumo = [], [], []

    # ---- pré-passo: classifica cada parágrafo (permite olhar "para frente") ----
    P = []
    n_eq = 0
    for para, em_tab in _iterar_paragrafos(doc):
        estilo = para.style.name if para.style is not None else ""
        rot = numeracao.rotulo(para)
        segs, eqs = _segmentos_do_paragrafo(para)
        n_eq += eqs
        heading = bool(RE_HEADING.match(estilo)) and not em_tab
        if rot and not heading:
            segs.insert(0, _novo_seg(rot + " "))
        plano = "".join(s["t"] for s in segs)
        P.append({
            "para": para, "segs": segs, "plano": plano, "heading": heading,
            "mq_forte": None if heading else RE_Q_FORTE.match(plano),
            "mq_fraca": None if heading else RE_Q_FRACA.match(plano),
            "malt": None if heading else RE_ALT.match(plano),
            "imgs": _extrair_imagens_do_paragrafo(para, doc_part),
        })

    def ha_alternativa_adiante(i, janela=6):
        for j in range(i + 1, min(i + 1 + janela, len(P))):
            q = P[j]
            if q["heading"] or q["mq_forte"] or q["mq_fraca"]:
                return False
            if q["malt"]:
                return True
        return False

    latex = PREAMBULO_BASE
    dentro_enum = False
    alt_aberta = False
    ultima_letra = None
    imagens = []
    ultima_q = None
    alternativas_iniciadas = False
    ultimo_rejeitado = False
    disc_nome, q_atual = "Geral", None

    def fechar_enum():
        nonlocal dentro_enum, latex, alt_aberta, ultima_letra
        if dentro_enum:
            latex += "\\end{enumerate}\n\n"
            dentro_enum = False
        alt_aberta = False
        ultima_letra = None

    for i, p in enumerate(P):
        if p["heading"]:
            fechar_enum()
            nome = re.sub(r"^\s*disciplina\s*:\s*", "", p["plano"].strip(), flags=re.IGNORECASE)
            nome_tex = _segmentos_para_latex([_novo_seg(nome)])
            latex += "\n% ==========================================\n"
            latex += f"\\section*{{DISCIPLINA: {nome_tex}}}\n"
            latex += "% ==========================================\n"
            ultima_q, alternativas_iniciadas, ultimo_rejeitado = None, False, False
            disc_nome, q_atual = nome or "Geral", None
            continue

        plano, segs = p["plano"], p["segs"]
        texto_vazio = not plano.strip()
        mq_f, mq_w, malt = p["mq_forte"], p["mq_fraca"], p["malt"]

        if not texto_vazio:
            mq = mq_f or mq_w
            numero = int(mq.group(1)) if mq else None
            esperado = (ultima_q + 1) if ultima_q is not None else None
            aceitar_q = bool(mq_f) or (bool(mq_w) and (
                ultima_q is None or alternativas_iniciadas
                or (not ultimo_rejeitado and numero == esperado)))

            if aceitar_q:
                fechar_enum()
                if ultima_q is not None and numero != ultima_q + 1:
                    avisos.append(f"Numeração fora de sequência em '{disc_nome}': Questão {numero} veio depois da {ultima_q}.")
                resto = _segmentos_para_latex(_cortar_prefixo(segs, mq.end()))
                latex += f"\\subsection*{{Questão {numero}}}\n"
                if resto:
                    latex += f"{resto}\n\n"
                q_atual = {"Disciplina": disc_nome, "Questão": numero, "Alternativas": 0}
                resumo.append(q_atual)
                ultima_q, alternativas_iniciadas, ultimo_rejeitado = numero, False, False

            elif malt and (dentro_enum or (malt.group(1) or malt.group(2)).lower() == "a"):
                letra = (malt.group(1) or malt.group(2))
                if dentro_enum and letra.lower() == "a" and (ultima_letra or "a").lower() != "a":
                    fechar_enum()  # lista recomeçou em "a": é outra lista
                if not dentro_enum:
                    latex += f"\\begin{{enumerate}}[{_bracket_enum(malt)}]\n"
                    dentro_enum = True
                resto = _segmentos_para_latex(_cortar_prefixo(segs, malt.end()))
                latex += f"\\item {resto}\n".replace("\\item \n", "\\item\n")
                alt_aberta, ultima_letra = True, letra
                alternativas_iniciadas, ultimo_rejeitado = True, False
                if q_atual is not None:
                    q_atual["Alternativas"] += 1

            elif dentro_enum and alt_aberta and ha_alternativa_adiante(i):
                latex += _segmentos_para_latex(segs) + "\n"  # continuação da alternativa (fórmula, 2ª linha)
                ultimo_rejeitado = False

            else:
                if dentro_enum:
                    notas.append(
                        f"Texto logo após as alternativas ({disc_nome}, Questão {ultima_q}) foi tratado como texto "
                        f"fora da questão: «{plano.strip()[:50]}…». Se era continuação da última alternativa, ajuste no .tex.")
                fechar_enum()
                latex += _segmentos_para_latex(segs) + "\n\n"
                ultimo_rejeitado = bool(mq_w) and not mq_f

        # ---- imagens: DEPOIS do texto do parágrafo, para ficarem na questão certa ----
        for blob, ext in p["imgs"]:
            idx = len(imagens) + 1
            latex += "\\begin{center}\n"
            latex += f"    \\includegraphics[width=0.6\\linewidth]{{images/image{idx}}}\n"
            latex += "\\end{center}\n"
            imagens.append((blob, ext))
            if ext not in ("png", "jpg", "pdf"):
                avisos.append(f"images/image{idx}.{ext}: formato '{ext}' não é aceito pelo pdfLaTeX — converta para PNG (mantendo o nome).")

    fechar_enum()
    latex += "\\end{document}"

    # ---- avisos globais ----
    if n_eq:
        avisos.append(f"{n_eq} equação(ões) do Word foram convertidas automaticamente para LaTeX — confira no PDF.")
    n_tab = len(doc.tables)
    if n_tab:
        avisos.append(f"O documento tem {n_tab} tabela(s): o texto foi lido célula a célula, como texto corrido. Tabelas de dados precisam ser refeitas em LaTeX.")
    if any(True for _ in doc.element.body.iter(qn("w:txbxContent"))):
        avisos.append("Há caixas de texto no documento: o conteúdo delas NÃO é lido pelo conversor.")
    for r in resumo:
        if r["Alternativas"] == 0:
            avisos.append(f"{r['Disciplina']} — Questão {r['Questão']}: nenhuma alternativa (a, b, c…) detectada (discursiva ou alternativas sem letra?).")
        elif r["Alternativas"] == 1:
            avisos.append(f"{r['Disciplina']} — Questão {r['Questão']}: só 1 alternativa detectada.")
    return latex, imagens, avisos, notas, resumo


def processar_acelerador_zip(docx_bytes):
    docx_bytes.seek(0)
    latex_text, imagens, avisos, notas, resumo = converter_docx_para_latex(docx_bytes)
    zip_buffer = BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as out_zip:
        out_zip.writestr("base.tex", latex_text)
        for idx, (blob, ext) in enumerate(imagens, start=1):
            out_zip.writestr(f"images/image{idx}.{ext}", blob)
    zip_buffer.seek(0)
    return zip_buffer, latex_text, avisos, notas, resumo


# ==========================================
# EMBARALHADOR (LATEX -> PROVAS)
# ==========================================

LETRAS = string.ascii_uppercase


class QuestaoObj:
    def __init__(self, uid, titulo_raw, titulo, corpo):
        self.uid = uid
        self.titulo_raw = titulo_raw      # conteúdo original das chaves do \subsection*
        self.titulo = titulo              # versão em texto plano (para CSV/avisos)
        self.corpo = corpo
        self.alternativas = []
        self.alt_prefacio = ""
        self.gabarito_orig = -1
        self.estilo_alternativa = "(a)"
        self.fixas = set()                # índices que ficam no fim ("todas as anteriores")


class BlocoObj:
    def __init__(self):
        self.texto_apoio = ""
        self.questoes = []


class DisciplinaObj:
    def __init__(self, nome, tag=None):
        self.nome = nome
        self.tag = tag      # \section*{...} original, reproduzido literalmente
        self.intro = ""     # texto entre o título da disciplina e a 1ª questão
        self.itens = []


class Prova:
    def __init__(self):
        self.preambulo = ""
        self.cabecalho = ""
        self.disciplinas = []
        self.rodape = r"\end{document}"
        self.avisos = []
        self.notas = []


_SECAO = r"\\(?:sub)?section\*?\{(?:[^{}]|\{[^{}]*\})*\}"
_INI = r"%[ \t]*IN[IÍ]CIO(?:[ \t]+DE)?[ \t]+BLOCO[^\n]*"
_FIM = r"%[ \t]*FIM(?:[ \t]+DE)?[ \t]+BLOCO[^\n]*"
RE_TOKENS = re.compile(f"({_SECAO}|{_INI}|{_FIM})", re.IGNORECASE)
RE_INI = re.compile(_INI, re.IGNORECASE)
RE_FIM = re.compile(_FIM, re.IGNORECASE)
RE_ENV = re.compile(r"\\(begin|end)\{(enumerate|itemize)\}")
RE_CORRETO = re.compile(r"(?<!\\)%[ \t]*(?:CORRET[OA]|CERTA)\b[^\n]*", re.IGNORECASE)
RE_FIXA = re.compile(
    r"\b(todas|nenhuma|ambas|qualquer)\b[^\n]{0,30}?\b(anteriores|alternativas|acima|op[cç][õo]es|afirmativas|"
    r"afirma[cç][õo]es|corretas?|incorretas?|erradas?|certas?|verdadeiras?|falsas?)\b", re.IGNORECASE)
RE_FIXA_SOLTA = re.compile(r"^\W*(todas|nenhuma|ambas)(\s+(elas|delas|as\s+duas|as\s+tr[êe]s))?\W*$", re.IGNORECASE)
RE_REF_QUESTAO = re.compile(r"\bquest(?:[ãa]o|[õo]es)\s+\d+", re.IGNORECASE)


def _plano_titulo(raw):
    return re.sub(r"\\[a-zA-Z]+\*?|[{}]", "", raw).strip()


def _fechar_chave(s, pos_abre):
    """Índice da '}' que fecha a '{' em pos_abre (ignora \\{ e \\}); -1 se não houver."""
    depth = 0
    i = pos_abre
    while i < len(s):
        c = s[i]
        if c == "\\":
            i += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return i
        i += 1
    return -1


def remover_hl(s):
    """Remove \\hl{...} (pacote soul) mantendo o conteúdo, mesmo com chaves aninhadas."""
    out, i = [], 0
    rx = re.compile(r"\\hl\s*\{")
    while True:
        m = rx.search(s, i)
        if not m:
            out.append(s[i:])
            break
        out.append(s[i:m.start()])
        fim = _fechar_chave(s, m.end() - 1)
        if fim == -1:
            i = m.end()
            continue
        out.append(remover_hl(s[m.end():fim]))
        i = fim + 1
    return "".join(out)


def _envs_toplevel(texto):
    envs, depth, ini = [], 0, None
    for m in RE_ENV.finditer(texto):
        if m.group(1) == "begin":
            if depth == 0:
                ini = m
            depth += 1
        else:
            if depth == 0:
                continue
            depth -= 1
            if depth == 0 and ini is not None and ini.group(2) == m.group(2):
                envs.append((ini, m))
                ini = None
    return envs


def _dividir_itens(inner):
    """Divide o miolo de um enumerate em itens (ignora \\item de listas aninhadas)."""
    pos, depth = [], 0
    for m in re.finditer(r"\\item(?![a-zA-Z])|\\(begin|end)\{(?:enumerate|itemize)\}", inner):
        if m.group(0).startswith("\\item"):
            if depth == 0:
                pos.append(m)
        elif m.group(1) == "begin":
            depth += 1
        else:
            depth = max(0, depth - 1)
    itens = []
    for k, m in enumerate(pos):
        fim = pos[k + 1].start() if k + 1 < len(pos) else len(inner)
        itens.append(inner[m.end():fim])
    prefacio = inner[:pos[0].start()] if pos else inner
    return itens, prefacio


def _extrair_alternativas(q, disc_nome, avisos, notas):
    rotulo = f"{disc_nome} — {q.titulo}"
    envs = [(b, e) for b, e in _envs_toplevel(q.corpo) if b.group(2) == "enumerate"]
    if not envs:
        notas.append(f"{rotulo}: sem alternativas (tratada como discursiva; gabarito '—').")
        return
    b, e = envs[-1]
    miolo = q.corpo[b.end():e.start()]
    mb = re.match(r"\s*\[([^\]]*)\]", miolo)
    bracket = None
    if mb:
        bracket, miolo = mb.group(1), miolo[mb.end():]
    if bracket is not None and not re.search(r"[aA]|alph", bracket):
        avisos.append(f"{rotulo}: o último enumerate tem rótulo '[{bracket}]', que não parece de alternativas (a, b, c). "
                      "Ele NÃO foi embaralhado — confira.")
        return
    itens, prefacio = _dividir_itens(miolo)
    if bracket:
        q.estilo_alternativa = bracket

    corretos, hl_em = [], []
    limpas = []
    for idx, it in enumerate(itens):
        t = it.strip()
        if RE_CORRETO.search(t):
            corretos.append(idx)
            t = RE_CORRETO.sub("", t).strip()
        if re.search(r"\\hl\s*\{", t):
            hl_em.append(idx)
            t = remover_hl(t)
        limpas.append(t)

    q.alternativas = limpas
    q.alt_prefacio = prefacio.strip()
    q.fixas = {i for i, t in enumerate(limpas) if RE_FIXA.search(t) or RE_FIXA_SOLTA.match(t)}
    q.gabarito_orig = corretos[0] if corretos else -1
    q.corpo = q.corpo[:b.start()] + "[[ALTS]]" + q.corpo[e.end():]

    if not corretos:
        extra = f" (a alternativa {LETRAS[hl_em[0]]} tem \\hl{{}} — faltou o %CORRETO?)" if len(hl_em) == 1 else ""
        avisos.append(f"{rotulo}: sem %CORRETO identificado{extra}.")
    elif len(corretos) > 1:
        avisos.append(f"{rotulo}: {len(corretos)} alternativas marcadas %CORRETO — usei a primeira.")
    if len(limpas) == 1:
        avisos.append(f"{rotulo}: só 1 alternativa encontrada.")
    if len(limpas) > 26:
        avisos.append(f"{rotulo}: mais de 26 alternativas — o gabarito não é representável por letra.")
    if q.fixas:
        notas.append(f"{rotulo}: alternativa(s) {', '.join(LETRAS[i] for i in sorted(q.fixas) if i < 26)} "
                     "('todas/nenhuma das anteriores') fixada(s) no fim.")


def parse_latex_para_objetos(latex_content):
    """Levanta ValueError se a estrutura básica não for encontrada."""
    if r"\begin{document}" not in latex_content:
        raise ValueError("não encontrei \\begin{document} no texto.")
    pre, resto = latex_content.split(r"\begin{document}", 1)
    if r"\end{document}" not in resto:
        raise ValueError("não encontrei \\end{document} no texto.")
    corpo = resto.split(r"\end{document}")[0]

    prova = Prova()
    prova.preambulo = pre + r"\begin{document}"

    tokens = RE_TOKENS.split(corpo)
    if (len(tokens) - 1) % 2 != 0:
        tokens.append("")
    prova.cabecalho = tokens[0]

    disc_atual = DisciplinaObj("Geral")
    prova.disciplinas.append(disc_atual)
    bloco, questao = None, None
    pendente = ""
    uid = 0

    for i in range(1, len(tokens), 2):
        tag = tokens[i].strip()
        conteudo = tokens[i + 1]

        if RE_INI.fullmatch(tag):
            if bloco is not None:
                prova.avisos.append(f"{disc_atual.nome}: '% INICIO BLOCO' aberto sem '% FIM BLOCO' antes do próximo bloco.")
            bloco = BlocoObj()
            bloco.texto_apoio = conteudo
            disc_atual.itens.append(bloco)
            questao = None

        elif RE_FIM.fullmatch(tag):
            if bloco is None:
                prova.avisos.append(f"{disc_atual.nome}: '% FIM BLOCO' sem '% INICIO BLOCO' correspondente.")
            bloco, questao = None, None
            pendente = conteudo
            if conteudo.strip():
                prova.avisos.append(f"{disc_atual.nome}: há texto entre '% FIM BLOCO' e a próxima marcação — ele viaja junto com a próxima questão.")

        else:  # \section* ou \subsection*
            inner = re.sub(r"^\\(?:sub)?section\*?\{", "", tag)[:-1]
            plano = _plano_titulo(inner)

            if re.match(r"^disciplina\s*:", plano, re.IGNORECASE):
                if bloco is not None:
                    prova.avisos.append(f"{disc_atual.nome}: '% INICIO BLOCO' sem FIM antes da disciplina '{plano}'.")
                nome = re.sub(r"^disciplina\s*:\s*", "", plano, flags=re.IGNORECASE)
                disc_atual = DisciplinaObj(nome, tag)
                disc_atual.intro = conteudo
                prova.disciplinas.append(disc_atual)
                bloco, questao, pendente = None, None, ""

            elif re.match(r"^quest[aã]o\b", plano, re.IGNORECASE):
                uid += 1
                q = QuestaoObj(uid, inner, plano, pendente + conteudo)
                pendente = ""
                questao = q
                (bloco.questoes if bloco is not None else disc_atual.itens).append(q)

            else:  # outra seção qualquer: acompanha o contexto
                extra = f"\n{tag}\n{conteudo}"
                if questao is not None:
                    questao.corpo += extra
                elif bloco is not None:
                    bloco.texto_apoio += extra
                elif not disc_atual.itens:
                    disc_atual.intro += extra
                else:
                    pendente += extra

    if bloco is not None:
        prova.avisos.append(f"{disc_atual.nome}: '% INICIO BLOCO' aberto sem '% FIM BLOCO' no fim do documento.")

    for disc in prova.disciplinas:
        for item in disc.itens:
            if isinstance(item, BlocoObj):
                for m in RE_REF_QUESTAO.finditer(item.texto_apoio):
                    prova.avisos.append(f"{disc.nome}: texto de apoio cita «{m.group(0)}» — a numeração muda nas versões embaralhadas.")
                    break
                qs = item.questoes
            else:
                qs = [item]
            for q in qs:
                _extrair_alternativas(q, disc.nome, prova.avisos, prova.notas)
                m = RE_REF_QUESTAO.search(q.corpo)
                if m:
                    prova.avisos.append(f"{disc.nome} — {q.titulo}: cita «{m.group(0)}» no texto — a numeração muda nas versões embaralhadas.")

    prova.disciplinas = [d for d in prova.disciplinas if d.itens or d.tag]
    if not any(d.itens for d in prova.disciplinas):
        raise ValueError("não encontrei nenhuma questão (\\subsection*{Questão N}).")
    return prova


def _embaralhar(rng, lista, evitar_identidade=True):
    base = list(lista)
    if len(base) < 2:
        return base
    cand = base
    for _ in range(25):
        cand = base[:]
        rng.shuffle(cand)
        if cand != base or not evitar_identidade:
            return cand
    return cand


def _questoes_de(item):
    return item.questoes if isinstance(item, BlocoObj) else [item]


def gerar_latex_embaralhado(prova, seed, sufixo="B", emb_questoes=True, emb_alternativas=True,
                            reiniciar_por_disciplina=False):
    rng = random.Random(seed)
    out = [prova.preambulo, "\n", prova.cabecalho.strip("\n"), "\n"]
    gabarito = []
    contador = 1

    for disc in prova.disciplinas:
        if reiniciar_por_disciplina:
            contador = 1
        if disc.tag:
            out.append(f"\n{disc.tag}\n{disc.intro}\n")
        itens = _embaralhar(rng, disc.itens) if emb_questoes else list(disc.itens)

        for item in itens:
            if isinstance(item, BlocoObj):
                out.append(f"\n{item.texto_apoio}\n")

            for q in _questoes_de(item):
                titulo_novo = re.sub(r"\d+", str(contador), q.titulo_raw, count=1)
                if not re.search(r"\d", q.titulo_raw):
                    titulo_novo = f"Questão {contador}"
                out.append(f"\\subsection*{{{titulo_novo}}}\n")

                n = len(q.alternativas)
                if n == 0:
                    nova_letra = "—"
                    corpo = q.corpo.replace("[[ALTS]]", "")
                else:
                    idx = list(range(n))
                    livres = [k for k in idx if k not in q.fixas]
                    fixas = [k for k in idx if k in q.fixas]
                    if emb_alternativas:
                        livres = _embaralhar(rng, livres)
                    ordem = livres + fixas
                    bloco = f"\\begin{{enumerate}}[{q.estilo_alternativa}]\n"
                    if q.alt_prefacio:
                        bloco += q.alt_prefacio + "\n"
                    nova_letra = "?"
                    for novo_i, orig_i in enumerate(ordem):
                        bloco += f"\\item {q.alternativas[orig_i]}\n"
                        if orig_i == q.gabarito_orig and novo_i < 26:
                            nova_letra = LETRAS[novo_i]
                    bloco += "\\end{enumerate}\n"
                    corpo = q.corpo.replace("[[ALTS]]", bloco) if "[[ALTS]]" in q.corpo else q.corpo + "\n" + bloco
                out.append(corpo)

                gabarito.append({
                    "uid": q.uid, "Disciplina": disc.nome, "Questão Nova": contador,
                    "Gabarito": nova_letra, "Origem": q.titulo, "Versão": f"Prova {sufixo}",
                })
                contador += 1

    out.append(prova.rodape)
    return "".join(out), gabarito


def gerar_versoes(prova, seed, n_versoes=2, **opcoes):
    """Gera B, C, D... garantindo que nenhuma seja idêntica a outra (quando possível)."""
    letras = LETRAS[1:1 + n_versoes]
    resultados, vistos = {}, set()
    for k, suf in enumerate(letras):
        tentativa = 0
        while True:
            tex, gab = gerar_latex_embaralhado(prova, seed + 10 * k + tentativa * 1000, suf, **opcoes)
            if tex not in vistos or tentativa >= 20:
                break
            tentativa += 1
        vistos.add(tex)
        resultados[suf] = (tex, gab)
    return resultados


def gabarito_mestre(prova, resultados):
    linhas = []
    mapas = {suf: {r["uid"]: r for r in gab} for suf, (_, gab) in resultados.items()}
    for disc in prova.disciplinas:
        for item in disc.itens:
            for q in _questoes_de(item):
                if not q.alternativas:
                    orig = "—"
                elif 0 <= q.gabarito_orig < 26:
                    orig = LETRAS[q.gabarito_orig]
                else:
                    orig = "?"
                linha = {"Disciplina": disc.nome, "Questão (A)": q.titulo, "Gab. A": orig}
                for suf, mp in mapas.items():
                    linha[f"Nº {suf}"] = mp[q.uid]["Questão Nova"]
                    linha[f"Gab. {suf}"] = mp[q.uid]["Gabarito"]
                linhas.append(linha)
    return pd.DataFrame(linhas)


def _csv(df):
    return df.to_csv(index=False, sep=";", encoding="utf-8-sig")


# ==========================================
# INTERFACE
# ==========================================

def _mostrar_avisos(avisos, notas, titulo_aviso="⚠️ Atenção antes de seguir"):
    if avisos:
        corpo = "\n".join(f"- {a}" for a in avisos[:15])
        if len(avisos) > 15:
            corpo += f"\n- … e mais {len(avisos) - 15} aviso(s) (veja a lista completa abaixo)."
        st.warning(f"**{titulo_aviso}**\n\n{corpo}")
        if len(avisos) > 15:
            with st.expander("Lista completa de avisos"):
                st.markdown("\n".join(f"- {a}" for a in avisos))
    if notas:
        with st.expander(f"ℹ️ {len(notas)} observação(ões) informativa(s)"):
            st.markdown("\n".join(f"- {n}" for n in notas))


def ui_acelerador():
    st.header("Conversor Inteligente de Word para LaTeX")
    st.markdown("""
    **Como usar:**
    1. Baixe o documento do Google Docs em `Arquivo > Fazer download > Microsoft Word (.docx)`.
    2. Suba o ficheiro abaixo.
    3. O sistema extrai **texto, equações, formatação básica e imagens** para você subir no Overleaf.
    """)
    file_docx = st.file_uploader("📂 Faça o upload da Prova em .docx", type=["docx"])
    if not file_docx:
        return
    chave = (file_docx.name, file_docx.size)

    if st.button("⚙️ Processar e Extrair Imagens", type="primary"):
        try:
            zip_buffer, preview, avisos, notas, resumo = processar_acelerador_zip(BytesIO(file_docx.getvalue()))
            st.session_state["acel"] = {
                "chave": chave, "zip": zip_buffer.getvalue(), "preview": preview,
                "avisos": avisos, "notas": notas, "resumo": resumo,
            }
        except Exception as e:  # noqa: BLE001
            st.session_state.pop("acel", None)
            st.error(f"Ocorreu um erro ao processar: {e}")

    r = st.session_state.get("acel")
    if r and r["chave"] == chave:
        st.success(f"✅ Conversão concluída: {len(r['resumo'])} questão(ões) detectada(s).")
        _mostrar_avisos(r["avisos"], r["notas"])
        st.info("💡 Extraia o .zip no Overleaf. Antes do **Embaralhador**, abra o `base.tex` e adicione `%CORRETO` na alternativa certa "
                "e `% INICIO BLOCO` / `% FIM BLOCO` em volta de cada **texto de apoio + as questões que dependem dele**.")
        st.download_button("📥 Baixar Pacote Base (.zip com imagens)", data=r["zip"],
                           file_name="Prova_Base_LaTeX.zip", mime="application/zip", use_container_width=True)
        if r["resumo"]:
            with st.expander("🔎 Conferência: questões e nº de alternativas detectadas"):
                st.dataframe(pd.DataFrame(r["resumo"]), hide_index=True)
        with st.expander("👀 Ver Prévia do Código Gerado"):
            st.code(r["preview"], language="latex")


def ui_embaralhador():
    st.header("Gerador de Versões (B, C, ...)")
    with st.expander("📖 INSTRUÇÕES PARA O EDITOR (Clique para expandir)", expanded=True):
        st.markdown("""
        ### Antes de gerar, a **Prova A** no Overleaf precisa ter:
        * `\\section*{DISCIPLINA: Nome}` — um por disciplina;
        * `%CORRETO` na alternativa certa de cada questão (a marcação `\\hl{}` é removida das versões automaticamente);
        * `% INICIO BLOCO` (ou `% INÍCIO BLOCO`) **antes do texto de apoio** e `% FIM BLOCO` **depois da última questão que usa esse texto**
          — o bloco inteiro é embaralhado como uma unidade.

        ### Depois
        1. Suba o `.tex` (ou cole o código) e clique em Gerar.
        2. Baixe o `.zip`, extraia e suba os `main_B.tex`, `main_C.tex`… na mesma pasta da Prova A no Overleaf.
        3. O `Gabarito_Mestre.csv` traz a Prova A e todas as versões lado a lado.
        """)

    tex_upload = st.file_uploader("📂 (Opcional) Suba o ficheiro .tex em vez de colar", type=["tex"])
    if tex_upload:
        bruto = tex_upload.getvalue()
        try:
            texto = bruto.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = bruto.decode("latin-1")
        st.caption(f"Usando o arquivo **{tex_upload.name}** (o campo de texto abaixo é ignorado enquanto houver arquivo).")
        latex_input = texto
    else:
        latex_input = st.text_area("Cole o Código LaTeX da Prova A Original aqui:", height=300)

    c1, c2, c3 = st.columns([1, 1, 2])
    with c1:
        seed_val = int(st.number_input("Semente inicial (seed)", value=42, step=1))
    with c2:
        n_versoes = int(st.number_input("Nº de versões extras", min_value=1, max_value=5, value=2, step=1))
    with c3:
        emb_q = st.checkbox("Embaralhar a ordem das questões", value=True)
        emb_a = st.checkbox("Embaralhar as alternativas", value=True)
        reiniciar = st.checkbox("Reiniciar a numeração a cada disciplina", value=False)

    opcoes = dict(emb_questoes=emb_q, emb_alternativas=emb_a, reiniciar_por_disciplina=reiniciar)
    chave = hashlib.md5(f"{latex_input}|{seed_val}|{n_versoes}|{sorted(opcoes.items())}".encode()).hexdigest()

    if st.button("🎲 Gerar versões (.zip)", type="primary"):
        if not latex_input.strip():
            st.warning("⚠️ Cole o código LaTeX (ou suba o .tex) antes de gerar.")
        else:
            try:
                prova = parse_latex_para_objetos(latex_input)
                resultados = gerar_versoes(prova, seed_val, n_versoes, **opcoes)
                mestre = gabarito_mestre(prova, resultados)
                zip_buffer = BytesIO()
                with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                    for suf, (tex, gab) in resultados.items():
                        df = pd.DataFrame(gab).drop(columns=["uid"])
                        zf.writestr(f"Prova_{suf}/main_{suf}.tex", tex)
                        zf.writestr(f"Prova_{suf}/Gabarito_{suf}.csv", _csv(df))
                    zf.writestr("Gabarito_Mestre.csv", _csv(mestre))
                st.session_state["emb"] = {
                    "chave": chave, "zip": zip_buffer.getvalue(), "mestre": mestre,
                    "avisos": prova.avisos, "notas": prova.notas,
                    "texs": {s: t for s, (t, _) in resultados.items()},
                }
            except ValueError as e:
                st.session_state.pop("emb", None)
                st.error(f"❌ Não foi possível ler a estrutura da prova: {e}")
            except Exception as e:  # noqa: BLE001
                st.session_state.pop("emb", None)
                st.error(f"❌ Erro inesperado ao gerar as provas: {e}")

    r = st.session_state.get("emb")
    if r and r["chave"] == chave:
        _mostrar_avisos(r["avisos"], r["notas"], "⚠️ Atenção antes de imprimir")
        st.success(f"✅ {len(r['texs'])} versão(ões) gerada(s).")
        st.download_button("📥 Baixar Pacote Completo de Embaralhamento (.zip)", data=r["zip"],
                           file_name=f"Provas_Embaralhadas_Seed{seed_val}.zip", mime="application/zip",
                           use_container_width=True)
        st.caption("🔍 Gabarito mestre (Prova A e versões):")
        st.dataframe(r["mestre"], hide_index=True)
        with st.expander("👀 Ver LaTeX de uma versão"):
            suf = st.selectbox("Versão", list(r["texs"]))
            st.code(r["texs"][suf], language="latex")


def main():
    st.set_page_config(page_title="ProvaOps - Sistema de Provas", layout="wide", page_icon="📝")
    st.title("🚀 Escola Analítica: ProvaOps")
    tab_a, tab_e = st.tabs(["⚡ 1. Acelerador (Extrair do Word)", "🎲 2. Embaralhador (Gerar Versões)"])
    with tab_a:
        ui_acelerador()
    with tab_e:
        ui_embaralhador()


if __name__ == "__main__":
    main()
